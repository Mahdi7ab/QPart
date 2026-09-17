from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION


TEMPLATE = r"C:\Users\m.abdollahi\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-strategy-memorandum\assets\reference.docx"
OUTPUT = r"D:\Mahdi\MyWorks\QPart\صورتجلسه جلسه دوم.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(table, color="D9D9D9", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_rtl(paragraph, rtl=True):
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1" if rtl else "0")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def set_run_font(run, name="Tahoma", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def clear_document_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Tahoma"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Tahoma")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Tahoma")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Tahoma")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color in (("Title", 20, "000000"), ("Heading 1", 14, "000000"), ("Heading 2", 11.5, "000000"), ("Heading 3", 10.5, "000000")):
        st = doc.styles[style_name]
        st.font.name = "Tahoma"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Tahoma")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Tahoma")
        st._element.rPr.rFonts.set(qn("w:cs"), "Tahoma")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 7)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.text = "صورتجلسه جلسه دوم"
    set_rtl(fp, False)
    for r in fp.runs:
        set_run_font(r, size=8.5, color="666666")


def add_paragraph(doc, text="", style="Normal", bold_prefix=None, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    set_rtl(p, align == WD_ALIGN_PARAGRAPH.RIGHT)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=10.5, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=10.5)
    else:
        r = p.add_run(text)
        if style == "Title":
            set_run_font(r, size=20, bold=True)
        elif style == "Heading 1":
            set_run_font(r, size=14, bold=True)
        elif style == "Heading 2":
            set_run_font(r, size=11.5, bold=True)
        else:
            set_run_font(r, size=10.5)
    return p


def format_table(table, header=True, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_cell_borders(table)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    for ri, row in enumerate(table.rows):
        if ri == 0 and header:
            set_repeat_table_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, "1F4E78" if ri == 0 and header else ("F4F7FA" if ri % 2 == 0 else "FFFFFF"))
            for p in cell.paragraphs:
                set_rtl(p, True)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    set_run_font(r, size=9.2, bold=(ri == 0 and header), color=("FFFFFF" if ri == 0 and header else "000000"))


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    format_table(table, widths=widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    return add_paragraph(doc, text, style=f"Heading {level}")


def main():
    doc = Document(TEMPLATE)
    clear_document_body(doc)
    style_document(doc)

    add_paragraph(doc, "صورتجلسه جلسه دوم راه‌اندازی کسب‌وکار فروش آنلاین", style="Title")
    p = add_paragraph(doc, "بررسی مدل فروش، تأمین کالا، بازاریابی و اقدامات اجرایی", align=WD_ALIGN_PARAGRAPH.RIGHT)
    for r in p.runs:
        set_run_font(r, size=11, color="4F4F4F")

    add_table(doc, ["عنوان", "اطلاعات"], [
        ("شماره جلسه", "جلسه دوم"),
        ("تاریخ جلسه", "اعلام نشده"),
        ("حاضرین", "اعلام نشده"),
        ("وضعیت سند", "جمع‌بندی جلسه"),
    ], widths=[1.7, 5.3])

    add_heading(doc, "۱. هدف جلسه", 1)
    add_paragraph(doc, "هدف جلسه، بررسی مسیر عملی راه‌اندازی یک کسب‌وکار فروش آنلاین در حوزه فعالیت موجود، با اتکا به تجربه و ارتباطات فعلی، حداقل سرمایه اولیه و مدل تأمین کالا پس از ثبت سفارش بود. همچنین درباره حفظ محرمانگی هویت، نحوه جذب مشتری، زیرساخت‌های نرم‌افزاری و برنامه جلسه بعد گفت‌وگو شد.")

    add_heading(doc, "۲. خلاصه مباحث مطرح‌شده", 1)
    add_heading(doc, "۲.۱ مدل اجرایی کم‌سرمایه", 2)
    add_paragraph(doc, "با توجه به محدودیت سرمایه، مدل پیشنهادی بر راه‌اندازی سایت و کانال‌های فروش، استفاده از سرمایه انسانی و ارتباطات موجود، و پرهیز از ایجاد انبار بزرگ در شروع کار استوار شد. در این مدل، کالاها تا حد امکان پس از ثبت سفارش مشتری از منابع قابل اعتماد تأمین و ارسال می‌شوند.")

    add_heading(doc, "۲.۲ تأمین کالا و زمان تحویل", 2)
    add_paragraph(doc, "مهم‌ترین چالش اجرایی، پر کردن فاصله زمانی میان ثبت سفارش و تأمین کالا عنوان شد. راهکار پیشنهادی این است که زمان تحویل هر محصول از ابتدا در سایت مشخص شود؛ برای نمونه، برخی اقلام طی دو تا سه روز کاری و برخی اقلام پس از تکمیل سبد سفارش ارسال شوند. برای سفارش‌های چندقلمی نیز می‌توان کالاها را از چند منبع جمع‌آوری، بسته‌بندی و یکجا ارسال کرد.")

    add_heading(doc, "۲.۳ محصولات، مشتریان و بازاریابی", 2)
    add_paragraph(doc, "توافق شد شروع کار بر حوزه‌ای باشد که تجربه، شناخت محصول و مشتری در آن وجود دارد و در مرحله نخست از ورود گسترده به کالاهای ناشناخته پرهیز شود. فهرستی اولیه از حدود هفت تا ده قلم کالای پرفروش یا زودبازده باید تهیه شود. بخشی از اقلام بدون خرید اولیه و بر اساس سفارش مشتری تأمین می‌شوند و برای بخشی دیگر، فقط در صورت اطمینان از فروش سریع، سرمایه محدود و کوتاه‌مدت اختصاص می‌یابد.")
    add_paragraph(doc, "همچنین استفاده از فهرست مشتریان موجود، پیامک، صفحه مجازی، ارتباط مستقیم و ابزارهای خودکار برای جذب و پیگیری مشتری مطرح شد. تأکید جلسه بر این بود که بازاریابی باید از مشتریان و ارتباطات موجود آغاز شود و بعد از روشن شدن مسیر، توسعه پیدا کند.")

    add_heading(doc, "۲.۴ زیرساخت، ارتباطات و اتوماسیون", 2)
    add_paragraph(doc, "تکمیل سامانه مدیریت ارتباط با مشتری برای ثبت محصولات و مشتریان، ایجاد یک خط ارتباطی مستقل برای کسب‌وکار، بررسی امکان استفاده از واتساپ و تلگرام، و در مراحل بعدی تهیه ابزار حسابداری یا ERP از موضوعات مطرح‌شده بود. همچنین ایده استخراج خودکار فهرست مشتریان بالقوه از صفحات و برچسب‌های مرتبط در شبکه‌های اجتماعی و پاسخ‌گویی خودکار به پیام‌ها مطرح شد تا زمان صرف‌شده برای امور تکراری کاهش یابد.")

    add_heading(doc, "۳. جمع‌بندی و تصمیمات جلسه", 1)
    decisions = [
        ("۱", "راه‌اندازی فروش آنلاین در حوزه فعالیت فعلی به‌عنوان مسیر اصلی شروع کسب‌وکار دنبال شود."),
        ("۲", "در مرحله اول، مدل فروش مبتنی بر سفارش و تأمین سریع کالا اجرا شود و از خواب سرمایه در انبار جلوگیری شود."),
        ("۳", "زمان تحویل هر محصول متناسب با مسیر تأمین آن در سایت اعلام شود تا انتظار مشتری از ابتدا روشن باشد."),
        ("۴", "فهرست اولیه محصولات بر اقلامی متمرکز شود که فروش آن‌ها آسان‌تر، تقاضای آن‌ها روشن‌تر و گردش سرمایه آن‌ها سریع‌تر است."),
        ("۵", "از مشتریان و ارتباطات موجود برای شروع بازاریابی استفاده شود و توسعه کانال‌های دیجیتال و اتوماسیون به‌صورت مرحله‌ای انجام گیرد."),
        ("۶", "تا حد امکان، اطلاعات و هویت مستقیم افراد در فرایند عمومی فروش دیده نشود و کانال‌ها و حساب‌های کاری مستقل ایجاد شوند."),
        ("۷", "جلسات بعدی با برنامه مشخص، به‌صورت هفتگی یا حداکثر هر ده روز یک‌بار برگزار و دو روز پیش از جلسه نهایی شوند."),
    ]
    add_table(doc, ["ردیف", "تصمیم یا جمع‌بندی"], decisions, widths=[0.7, 6.3])

    add_heading(doc, "۴. اقدامات تا جلسه بعد", 1)
    actions = [
        ("۱", "تکمیل سامانه مدیریت ارتباط با مشتری و آماده‌سازی امکان ثبت محصولات و مشتریان", "مسئول راه‌اندازی سایت و زیرساخت", "تا جلسه بعد", "در حال انجام"),
        ("۲", "تهیه نسخه نمایشی صفحه اصلی سایت و بررسی نمونه‌های معرفی‌شده", "مسئول راه‌اندازی سایت", "تا جلسه بعد", "در حال انجام"),
        ("۳", "تهیه فهرست حدود هفت تا ده محصول اولویت‌دار و تعیین روش تأمین، زمان تحویل و نیاز احتمالی به سرمایه", "مسئول تأمین کالا با همکاری طرفین", "تا جلسه بعد", "اولویت بالا"),
        ("۴", "استخراج فهرست مشتریان موجود شامل نام، شماره تماس و شهر، ترجیحاً در قالب Excel یا PDF", "مسئول دسترسی به اطلاعات مشتریان", "تا جلسه بعد", "نیازمند بررسی خروجی سیستم"),
        ("۵", "طراحی مسیر اولیه بازاریابی از طریق پیامک، صفحه مجازی و ارتباط مستقیم با مشتریان موجود", "مسئول بازاریابی و کانال‌های دیجیتال", "تا جلسه بعد", "پیشنهاد اولیه"),
        ("۶", "بررسی ایجاد خط ارتباطی و حساب‌های مستقل کاری برای مدیریت پیام‌ها و دسترسی مشترک", "طرفین", "تا جلسه بعد", "نیازمند تصمیم اجرایی"),
        ("۷", "بررسی یک ابزار ساده حسابداری برای شروع و تعیین زمان مناسب حرکت به سمت ERP", "طرفین", "جلسه بعد", "فعلاً بدون خرید قطعی"),
        ("۸", "تعیین تاریخ قطعی جلسه بعد و ارسال آن حداقل دو روز پیش از برگزاری", "طرفین", "حداکثر تا ده روز آینده", "ضروری"),
    ]
    add_table(doc, ["ردیف", "اقدام", "مسئول", "زمان هدف", "وضعیت"], actions, widths=[0.5, 3.0, 1.6, 1.1, 0.8])

    add_heading(doc, "۵. ریسک‌ها و ملاحظات اجرایی", 1)
    risks = [
        ("تأخیر در تأمین کالا", "بالا", "اعلام زمان تحویل واقعی برای هر محصول و استفاده از چند منبع تأمین قابل اعتماد."),
        ("محدودیت سرمایه اولیه", "بالا", "شروع با فروش سفارشی و خرید محدود فقط برای اقلام زودبازده و قابل پیش‌بینی."),
        ("وابستگی به زمان و مشغله روزمره", "متوسط", "تقویم ثابت جلسات، تقسیم مسئولیت و خودکارسازی کارهای تکراری."),
        ("آشکار شدن ارتباط مستقیم افراد با کسب‌وکار", "متوسط", "ایجاد کانال‌ها و حساب‌های مستقل و مشخص‌کردن سطح دسترسی به اطلاعات."),
        ("فرسایشی‌شدن تصمیم‌گیری و شروع کار", "متوسط", "تعیین خروجی مشخص برای هر جلسه و بررسی منظم اقدامات انجام‌شده و انجام‌نشده."),
    ]
    add_table(doc, ["مورد", "سطح", "راهکار یا ملاحظه"], risks, widths=[2.0, 1.0, 4.0])

    add_heading(doc, "۶. دستور جلسه پیشنهادی برای جلسه بعد", 1)
    for item in [
        "بررسی نسخه نمایشی سایت و وضعیت سامانه مدیریت ارتباط با مشتری.",
        "مرور فهرست محصولات اولویت‌دار، قیمت، روش تأمین و زمان تحویل هر قلم.",
        "بررسی فهرست مشتریان و انتخاب گروه اولیه برای اطلاع‌رسانی.",
        "تصمیم‌گیری درباره خط ارتباطی، حساب‌های کاری و شیوه مدیریت پیام‌ها.",
        "ارزیابی اقدامات انجام‌شده، اقدامات انجام‌نشده و دلیل هر مورد.",
        "تعیین گام اجرایی بعدی و زمان‌بندی شروع فروش آزمایشی.",
    ]:
        p = add_paragraph(doc, "• " + item)
        p.paragraph_format.left_indent = Inches(0.15)

    add_heading(doc, "۷. موارد نیازمند تعیین تکلیف", 1)
    add_paragraph(doc, "تاریخ دقیق جلسه، اسامی حاضرین، نام و مشخصات رسمی کسب‌وکار، فهرست نهایی محصولات مرحله اول، زمان‌های تحویل قابل تعهد، شیوه ثبت قیمت و حاشیه سود، و ابزار حسابداری مناسب باید در ادامه مشخص شوند.")

    add_heading(doc, "خاتمه جلسه", 1)
    add_paragraph(doc, "با جمع‌بندی اقدامات فوق، جلسه خاتمه یافت و مقرر شد پیشرفت موارد در جلسه بعد بررسی شود.")

    doc.save(OUTPUT)
    print("created")


if __name__ == "__main__":
    main()
